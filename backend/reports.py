import io
from datetime import datetime, timezone
from pathlib import Path

from analysis import _l, admin_accounts, summarize

LOGO = str(Path(__file__).parent / "assets" / "logo_small.png")
BLUE, RED, NAVY, GREY = "004070", "B03010", "002040", "5B6770"
DIVISION = "A Division of YFI Enterprises, LLC"
CONFIDENTIAL = ("CONFIDENTIAL - This report contains sensitive system and security information prepared by "
                "Temprudence Technologies for the named client. Do not distribute without authorization.")
FOOTER = f"Confidential - Temprudence Technologies - {DIVISION}"


def _s(v):
    if v is None or v == "":
        return "-"
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, list):
        return ", ".join(_s(x) for x in v) or "-"
    return str(v)


def _d(v):
    if not v:
        return "-"
    try:
        return datetime.fromisoformat(str(v).replace("Z", "+00:00")).strftime("%Y-%m-%d %H:%M")
    except ValueError:
        return str(v)


def _flag_table(flags):
    if not flags:
        return [{"type": "p", "text": "No issues were flagged in this scan."}]
    return [{"type": "table", "headers": ["Severity", "Category", "Finding", "Detail"],
             "rows": [[f["severity"].upper(), f["category"], f["title"], f.get("detail") or "-"] for f in flags]}]


def build_computer_report(computer, scan, prepared_by):
    d = scan["data"]
    sy, os_, hw = d.get("system") or {}, d.get("os") or {}, d.get("hardware") or {}
    net, sp, us, sec = d.get("network") or {}, d.get("speed") or {}, d.get("users") or {}, d.get("security") or {}
    sm = summarize(d)
    b = []
    b += [{"type": "h1", "text": "1. Executive Summary"},
          {"type": "kv", "rows": [["Computer name", sy.get("computer_name")], ["Owner / signed-in user", computer.get("owner") or sm["current_user"]],
                                  ["Operating system", f"{_s(os_.get('caption'))} {_s(os_.get('display_version'))}"],
                                  ["Build", os_.get("build")], ["Manufacturer / model", f"{_s(sy.get('manufacturer'))} {_s(sy.get('model'))}"],
                                  ["Serial number", sy.get("serial")], ["Domain / workgroup", sy.get("domain")],
                                  ["Primary IP", sm["ip"]], ["Installed applications", sm["software_count"]],
                                  ["Enabled local administrators", sm["admin_count"]], ["Scan date", _d(scan["scanned_at"])],
                                  ["Scanner ran as Administrator", d.get("is_admin")]]},
          {"type": "h2", "text": "Flagged items"}] + _flag_table(scan.get("flags") or [])
    b += [{"type": "h1", "text": "2. Operating System"},
          {"type": "kv", "rows": [["Edition", os_.get("caption")], ["Version", os_.get("display_version")], ["Kernel version", os_.get("version")],
                                  ["Build", os_.get("build")], ["Architecture", os_.get("architecture")], ["Install date", _d(os_.get("install_date"))],
                                  ["Last boot", _d(os_.get("last_boot"))], ["Activation", os_.get("activation_status")],
                                  ["Registered owner", os_.get("registered_owner")], ["Language", os_.get("language")]]},
          {"type": "h2", "text": f"Installed updates ({len(_l(os_.get('hotfixes')))})"},
          {"type": "table", "headers": ["Update", "Description", "Installed"],
           "rows": [[h.get("id"), h.get("description"), _d(h.get("installed_on"))] for h in _l(os_.get("hotfixes"))]}]
    b += [{"type": "h1", "text": "3. Hardware"},
          {"type": "kv", "rows": [["Manufacturer", sy.get("manufacturer")], ["Model", sy.get("model")], ["System type", sy.get("system_type")],
                                  ["Serial number", sy.get("serial")], ["BIOS", f"{_s(sy.get('bios_vendor'))} {_s(sy.get('bios_version'))} ({_d(sy.get('bios_date'))})"],
                                  ["Total memory", f"{_s(hw.get('memory_total_gb'))} GB"]]},
          {"type": "h2", "text": "Processors"},
          {"type": "table", "headers": ["Processor", "Cores", "Logical", "Max clock (MHz)"],
           "rows": [[c.get("name"), c.get("cores"), c.get("logical"), c.get("max_clock_mhz")] for c in _l(hw.get("cpu"))]},
          {"type": "h2", "text": "Memory modules"},
          {"type": "table", "headers": ["Slot", "Capacity (GB)", "Speed (MHz)", "Manufacturer", "Part number"],
           "rows": [[m.get("slot"), m.get("capacity_gb"), m.get("speed_mhz"), m.get("manufacturer"), m.get("part_number")] for m in _l(hw.get("memory_modules"))]},
          {"type": "h2", "text": "Volumes"},
          {"type": "table", "headers": ["Drive", "Label", "File system", "Size (GB)", "Free (GB)", "Free %"],
           "rows": [[v.get("device"), v.get("label"), v.get("filesystem"), v.get("size_gb"), v.get("free_gb"),
                     f"{round((v.get('free_gb') or 0) / v['size_gb'] * 100)}%" if v.get("size_gb") else "-"] for v in _l(hw.get("volumes"))]},
          {"type": "h2", "text": "Physical disks"},
          {"type": "table", "headers": ["Model", "Interface", "Media", "Size (GB)"],
           "rows": [[p.get("model"), p.get("interface"), p.get("media_type"), p.get("size_gb")] for p in _l(hw.get("physical_disks"))]},
          {"type": "h2", "text": "Graphics"},
          {"type": "table", "headers": ["Adapter", "Driver", "VRAM (MB)", "Resolution"],
           "rows": [[g.get("name"), g.get("driver_version"), g.get("vram_mb"), g.get("resolution")] for g in _l(hw.get("gpu"))]}]
    if _l(hw.get("battery")):
        b += [{"type": "h2", "text": "Battery"},
              {"type": "table", "headers": ["Battery", "Charge %", "Status"],
               "rows": [[x.get("name"), x.get("charge_pct"), x.get("status")] for x in _l(hw.get("battery"))]}]
    wifi, proxy = net.get("wifi") or {}, net.get("proxy") or {}
    b += [{"type": "h1", "text": "4. Network"},
          {"type": "h2", "text": "Adapters"},
          {"type": "table", "headers": ["Adapter", "Status", "MAC", "IPv4 / subnet", "Gateway", "DNS", "DHCP", "Link"],
           "rows": [[a.get("name"), a.get("status"), a.get("mac"),
                     " ".join(f"{ip}/{sn}" for ip, sn in zip(_l(a.get("ipv4")), _l(a.get("subnet")) + [""] * 8)) or "-",
                     _s(a.get("gateway")), _s(a.get("dns")), a.get("dhcp"), a.get("link_speed")] for a in _l(net.get("adapters"))]},
          {"type": "h2", "text": "Throughput (measured to console)"},
          {"type": "kv", "rows": [["Latency (avg)", f"{_s(sp.get('latency_ms'))} ms"], ["Download", f"{_s(sp.get('download_mbps'))} Mbps"],
                                  ["Upload", f"{_s(sp.get('upload_mbps'))} Mbps"], ["Test target", sp.get("target")], ["Notes", sp.get("error")]]},
          {"type": "h2", "text": "Wi-Fi"},
          {"type": "kv", "rows": [[k.replace("_", " ").title(), v] for k, v in wifi.items()] or [["Wi-Fi", "Not connected / not present"]]},
          {"type": "h2", "text": "Firewall profiles"},
          {"type": "table", "headers": ["Profile", "Enabled", "Inbound", "Outbound"],
           "rows": [[f.get("profile"), f.get("enabled"), f.get("inbound"), f.get("outbound")] for f in _l(net.get("firewall"))]},
          {"type": "h2", "text": "Proxy"},
          {"type": "kv", "rows": [["User proxy enabled", proxy.get("enabled")], ["Proxy server", proxy.get("server")],
                                  ["Auto-config URL", proxy.get("auto_config_url")], ["WinHTTP", proxy.get("winhttp")]]}]
    b += [{"type": "h1", "text": "5. Users & Privileges"},
          {"type": "kv", "rows": [["Signed-in user", us.get("current_user") or sy.get("current_user")], ["Administrators group", us.get("admin_group")],
                                  ["Enabled local administrators", ", ".join(a["name"] for a in admin_accounts(d)) or "-"]]},
          {"type": "h2", "text": "Local accounts"},
          {"type": "table", "headers": ["Account", "Enabled", "Administrator", "Last logon", "Password required", "Groups"],
           "rows": [[u.get("name"), u.get("enabled"), u.get("is_admin"), _d(u.get("last_logon")), u.get("password_required"), _s(u.get("groups"))]
                    for u in _l(us.get("local"))]},
          {"type": "h2", "text": "Group memberships"},
          {"type": "table", "headers": ["Group", "Members"],
           "rows": [[g.get("name"), ", ".join(m.get("name", "") for m in _l(g.get("members")))] for g in _l(us.get("groups")) if _l(g.get("members"))]}]
    dfd = sec.get("defender") or {}
    b += [{"type": "h1", "text": "6. Security & Settings"},
          {"type": "kv", "rows": [["Time zone", sy.get("timezone")], ["Domain joined", sy.get("part_of_domain")], ["Domain / workgroup", sy.get("domain")],
                                  ["UAC enabled", sec.get("uac_enabled")], ["Secure Boot", sec.get("secure_boot")],
                                  ["TPM ready", (sec.get("tpm") or {}).get("ready")], ["Defender real-time", dfd.get("realtime")],
                                  ["Defender signature age (days)", dfd.get("signature_age_days")]]},
          {"type": "h2", "text": "Antivirus"},
          {"type": "table", "headers": ["Product", "Enabled", "Up to date"],
           "rows": [[a.get("name"), a.get("enabled"), a.get("up_to_date")] for a in _l(sec.get("antivirus"))]},
          {"type": "h2", "text": "BitLocker"},
          {"type": "table", "headers": ["Volume", "Protection", "Status", "Encrypted %", "Method"],
           "rows": [[x.get("mount"), x.get("protection"), x.get("status"), x.get("encryption_pct"), x.get("method")] for x in _l(sec.get("bitlocker"))]}]
    sw = sorted(_l(d.get("software")), key=lambda s: str(s.get("name", "")).lower())
    b += [{"type": "h1", "text": f"7. Installed Software ({len(sw)})"},
          {"type": "table", "headers": ["Application", "Version", "Publisher", "Installed", "Type"],
           "rows": [[s.get("name"), s.get("version"), s.get("publisher"), s.get("install_date"), s.get("arch")] for s in sw]}]
    return {"title": "Computer Inventory Report", "subtitle": sy.get("computer_name") or computer.get("name"),
            "meta": [["Computer", computer.get("name")], ["Serial", sy.get("serial")], ["Scan date", _d(scan["scanned_at"])],
                     ["Prepared by", prepared_by], ["Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")]],
            "blocks": b}


def build_fleet_report(computers, scans, prepared_by):
    rows, os_dist, sw_count, all_flags, outdated, admins, low = [], {}, {}, [], [], [], []
    for c in computers:
        s = scans.get(c["id"])
        if not s:
            continue
        d, sm, fl = s["data"], c.get("summary") or {}, s.get("flags") or []
        key = f"{_s(sm.get('os_caption'))} {_s(sm.get('os_version'))}"
        os_dist[key] = os_dist.get(key, 0) + 1
        for a in _l(d.get("software")):
            n = a.get("name")
            if n:
                e = sw_count.setdefault(n, {"count": 0, "versions": set()})
                e["count"] += 1
                e["versions"].add(str(a.get("version") or "-"))
        rows.append([c["name"], key, sm.get("os_build"), _d(c.get("last_scan_at")), sm.get("admin_count"), sm.get("software_count"),
                     f"{_s(sm.get('disk_free_gb'))} / {_s(sm.get('disk_total_gb'))} GB", (c.get("status") or "ok").upper()])
        for f in fl:
            all_flags.append([c["name"], f["severity"].upper(), f["category"], f["title"], f.get("detail") or "-"])
            if f["code"] in ("os_eol", "os_outdated"):
                outdated.append([c["name"], key, sm.get("os_build"), f["title"]])
            if f["code"] == "low_disk":
                low.append([c["name"], f["title"], f.get("detail")])
        admins.append([c["name"], len(admin_accounts(d)), ", ".join(a["name"] for a in admin_accounts(d)) or "-"])
    top = sorted(sw_count.items(), key=lambda kv: (-kv[1]["count"], kv[0].lower()))[:30]
    n = len(rows)
    b = [{"type": "h1", "text": "1. Fleet Overview"},
         {"type": "kv", "rows": [["Computers in report", n], ["Computers with critical findings", sum(1 for r in rows if r[7] == "CRITICAL")],
                                 ["Computers with warnings", sum(1 for r in rows if r[7] == "WARNING")], ["Total findings", len(all_flags)],
                                 ["Distinct applications", len(sw_count)]]},
         {"type": "h2", "text": "Operating system distribution"},
         {"type": "table", "headers": ["Operating system", "Computers", "Share"],
          "rows": [[k, v, f"{round(v / n * 100)}%"] for k, v in sorted(os_dist.items(), key=lambda x: -x[1])]},
         {"type": "h1", "text": "2. Computers"},
         {"type": "table", "headers": ["Computer", "OS", "Build", "Last scan", "Admins", "Apps", "Disk free", "Status"], "rows": rows},
         {"type": "h1", "text": "3. Outdated or Unsupported OS"}]
    b += [{"type": "table", "headers": ["Computer", "OS", "Build", "Finding"], "rows": outdated}] if outdated else [{"type": "p", "text": "All computers run a supported, current OS build."}]
    b += [{"type": "h1", "text": "4. Administrator Accounts"},
          {"type": "table", "headers": ["Computer", "Enabled local admins", "Accounts"], "rows": admins},
          {"type": "h1", "text": "5. Low Disk Space"}]
    b += [{"type": "table", "headers": ["Computer", "Volume", "Detail"], "rows": low}] if low else [{"type": "p", "text": "No volumes are low on free space."}]
    b += [{"type": "h1", "text": "6. Software"},
          {"type": "h2", "text": "Most common applications"},
          {"type": "table", "headers": ["Application", "Installs", "Versions"],
           "rows": [[k, v["count"], ", ".join(sorted(v["versions"]))[:120]] for k, v in top]},
          {"type": "h1", "text": "7. All Findings"}]
    b += [{"type": "table", "headers": ["Computer", "Severity", "Category", "Finding", "Detail"], "rows": all_flags}] if all_flags else [{"type": "p", "text": "No findings."}]
    return {"title": "Fleet Inventory Summary", "subtitle": f"{n} computer{'s' if n != 1 else ''}",
            "meta": [["Computers", n], ["Prepared by", prepared_by], ["Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")]],
            "blocks": b}


# ------------------------------- DOCX -------------------------------
def render_docx(rep):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.orientation = WD_ORIENT.PORTRAIT
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Inches(0.75))
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Aptos", Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    rgb = lambda h: RGBColor.from_string(h)

    def para(text, size=10, bold=False, color=None, align=None, space=4, italic=False):
        p = doc.add_paragraph()
        r = p.add_run(_s(text))
        r.font.size, r.bold, r.italic = Pt(size), bold, italic
        if color:
            r.font.color.rgb = rgb(color)
        if align:
            p.alignment = align
        p.paragraph_format.space_after = Pt(space)
        return p

    def shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hexcolor)
        tcPr.append(shd)

    def table(headers, rows, kv=False):
        if not rows:
            para("None recorded.", italic=True, color=GREY)
            return
        t = doc.add_table(rows=0, cols=len(headers) if headers else 2)
        t.style = "Table Grid"
        if headers:
            cells = t.add_row().cells
            for i, h in enumerate(headers):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(h)
                r.bold, r.font.size = True, Pt(9)
                r.font.color.rgb = rgb("FFFFFF")
                shade(cells[i], BLUE)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(_s(v))
                r.font.size = Pt(8.5 if not kv else 9.5)
                if kv and i == 0:
                    r.bold = True
                    r.font.color.rgb = rgb(NAVY)
                    shade(cells[i], "EEF3F8")
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    p.add_run().add_picture(LOGO, width=Inches(2.6))
    para(rep["title"], 26, True, BLUE, WD_ALIGN_PARAGRAPH.CENTER, 6)
    para(rep["subtitle"], 16, False, NAVY, WD_ALIGN_PARAGRAPH.CENTER, 24)
    table(None, rep["meta"], kv=True)
    para(CONFIDENTIAL, 9, True, RED, WD_ALIGN_PARAGRAPH.CENTER, 8)
    para(DIVISION, 10, False, GREY, WD_ALIGN_PARAGRAPH.CENTER)
    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

    hdr = sec.header.paragraphs[0]
    hr = hdr.add_run(f"TEMPRUDENCE TECHNOLOGIES  |  {rep['title']}  |  {_s(rep['subtitle'])}")
    hr.font.size, hr.font.color.rgb, hr.bold = Pt(8), rgb(BLUE), True
    ft = sec.footer.paragraphs[0]
    ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = ft.add_run(FOOTER)
    fr.font.size, fr.font.color.rgb = Pt(8), rgb(GREY)

    for blk in rep["blocks"]:
        t = blk["type"]
        if t == "h1":
            h = doc.add_heading(blk["text"], level=1)
            for r in h.runs:
                r.font.color.rgb, r.font.name, r.font.size = rgb(BLUE), "Aptos", Pt(16)
        elif t == "h2":
            h = doc.add_heading(blk["text"], level=2)
            for r in h.runs:
                r.font.color.rgb, r.font.name, r.font.size = rgb(NAVY), "Aptos", Pt(12)
        elif t == "p":
            para(blk["text"])
        elif t == "kv":
            table(None, blk["rows"], kv=True)
        elif t == "table":
            table(blk["headers"], blk["rows"])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------- ODT -------------------------------
def render_odt(rep):
    from odf.draw import Frame, Image as DImage
    from odf.opendocument import OpenDocumentText
    from odf.style import (Footer, Header, MasterPage, PageLayout, PageLayoutProperties, ParagraphProperties, Style,
                           TableCellProperties, TableColumnProperties, TableProperties, TextProperties, GraphicProperties)
    from odf.table import Table, TableCell, TableColumn, TableRow
    from odf.text import H, P

    doc = OpenDocumentText()
    pl = PageLayout(name="TTPage")
    pl.addElement(PageLayoutProperties(pagewidth="8.5in", pageheight="11in", margintop="0.6in", marginbottom="0.6in",
                                       marginleft="0.75in", marginright="0.75in", printorientation="portrait"))
    doc.automaticstyles.addElement(pl)

    def pstyle(name, size="10pt", bold=False, color=f"#{NAVY}", align=None, family="paragraph", before="0in", after="0.06in", brk=None):
        s = Style(name=name, family=family)
        pp = {"marginbottom": after, "margintop": before}
        if align:
            pp["textalign"] = align
        if brk:
            pp["breakbefore"] = brk
        s.addElement(ParagraphProperties(**pp))
        tp = {"fontsize": size, "color": color, "fontfamily": "Aptos", "fontname": "Aptos"}
        if bold:
            tp["fontweight"] = "bold"
        s.addElement(TextProperties(**tp))
        doc.styles.addElement(s)
        return s

    from odf.style import FontFace
    doc.fontfacedecls.addElement(FontFace(name="Aptos", fontfamily="Aptos, Arial", fontfamilygeneric="swiss"))
    s_norm = pstyle("TTBody", color="#1F2933")
    s_h1 = pstyle("TTH1", "16pt", True, f"#{BLUE}", before="0.18in", after="0.08in")
    s_h2 = pstyle("TTH2", "12pt", True, f"#{NAVY}", before="0.12in", after="0.05in")
    s_title = pstyle("TTTitle", "26pt", True, f"#{BLUE}", "center", before="0.1in")
    s_sub = pstyle("TTSub", "16pt", False, f"#{NAVY}", "center", after="0.3in")
    s_conf = pstyle("TTConf", "9pt", True, f"#{RED}", "center", before="0.2in")
    s_div = pstyle("TTDiv", "10pt", False, f"#{GREY}", "center")
    s_center = pstyle("TTCenter", align="center", before="0.4in")
    s_brk = pstyle("TTBreak", brk="page")
    s_cell = pstyle("TTCell", "8.5pt", color="#1F2933", after="0in")
    s_cellh = pstyle("TTCellH", "9pt", True, "#FFFFFF", after="0in")
    s_cellk = pstyle("TTCellK", "9pt", True, f"#{NAVY}", after="0in")
    s_foot = pstyle("TTFoot", "8pt", False, f"#{GREY}", "center")
    s_head = pstyle("TTHead", "8pt", True, f"#{BLUE}")

    mp = MasterPage(name="Standard", pagelayoutname=pl)
    hd = Header()
    hd.addElement(P(stylename=s_head, text=f"TEMPRUDENCE TECHNOLOGIES  |  {rep['title']}  |  {_s(rep['subtitle'])}"))
    mp.addElement(hd)
    ft = Footer()
    ft.addElement(P(stylename=s_foot, text=FOOTER))
    mp.addElement(ft)
    doc.masterstyles.addElement(mp)

    c_head = Style(name="TTCH", family="table-cell")
    c_head.addElement(TableCellProperties(backgroundcolor=f"#{BLUE}", border="0.5pt solid #9AA5B1", padding="0.04in"))
    c_body = Style(name="TTCB", family="table-cell")
    c_body.addElement(TableCellProperties(border="0.5pt solid #9AA5B1", padding="0.04in"))
    c_key = Style(name="TTCK", family="table-cell")
    c_key.addElement(TableCellProperties(backgroundcolor="#EEF3F8", border="0.5pt solid #9AA5B1", padding="0.04in"))
    t_style = Style(name="TTTable", family="table")
    t_style.addElement(TableProperties(width="7in", align="margins"))
    for s in (c_head, c_body, c_key, t_style):
        doc.automaticstyles.addElement(s)
    counter = [0]

    def table(headers, rows, kv=False):
        if not rows:
            doc.text.addElement(P(stylename=s_norm, text="None recorded."))
            return
        counter[0] += 1
        ncol = len(headers) if headers else 2
        t = Table(name=f"T{counter[0]}", stylename=t_style)
        t.addElement(TableColumn(numbercolumnsrepeated=ncol))
        if headers:
            tr = TableRow()
            for h in headers:
                tc = TableCell(stylename=c_head, valuetype="string")
                tc.addElement(P(stylename=s_cellh, text=h))
                tr.addElement(tc)
            t.addElement(tr)
        for row in rows:
            tr = TableRow()
            for i, v in enumerate(row):
                key = kv and i == 0
                tc = TableCell(stylename=c_key if key else c_body, valuetype="string")
                tc.addElement(P(stylename=s_cellk if key else s_cell, text=_s(v)))
                tr.addElement(tc)
            t.addElement(tr)
        doc.text.addElement(t)
        doc.text.addElement(P(stylename=s_norm, text=""))

    fr_style = Style(name="TTFrame", family="graphic")
    fr_style.addElement(GraphicProperties(anchortype="as-char"))
    doc.automaticstyles.addElement(fr_style)
    p = P(stylename=s_center)
    frame = Frame(stylename=fr_style, width="2.6in", height="2.6in", anchortype="as-char", name="logo")
    frame.addElement(DImage(href=doc.addPicture(LOGO)))
    p.addElement(frame)
    doc.text.addElement(p)
    doc.text.addElement(P(stylename=s_title, text=rep["title"]))
    doc.text.addElement(P(stylename=s_sub, text=_s(rep["subtitle"])))
    table(None, rep["meta"], kv=True)
    doc.text.addElement(P(stylename=s_conf, text=CONFIDENTIAL))
    doc.text.addElement(P(stylename=s_div, text=DIVISION))
    doc.text.addElement(P(stylename=s_brk, text=""))
    for blk in rep["blocks"]:
        t = blk["type"]
        if t == "h1":
            doc.text.addElement(H(outlinelevel=1, stylename=s_h1, text=blk["text"]))
        elif t == "h2":
            doc.text.addElement(H(outlinelevel=2, stylename=s_h2, text=blk["text"]))
        elif t == "p":
            doc.text.addElement(P(stylename=s_norm, text=blk["text"]))
        elif t == "kv":
            table(None, blk["rows"], kv=True)
        elif t == "table":
            table(blk["headers"], blk["rows"])
    buf = io.BytesIO()
    doc.write(buf)
    return buf.getvalue()


# ------------------------------- PDF -------------------------------
def render_pdf(rep):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    c = lambda h: colors.HexColor(f"#{h}")
    base = ParagraphStyle("b", fontName="Helvetica", fontSize=9.5, leading=12.5, textColor=c("1F2933"))
    cell = ParagraphStyle("c", parent=base, fontSize=7.8, leading=9.6)
    cellh = ParagraphStyle("ch", parent=cell, fontName="Helvetica-Bold", textColor=colors.white)
    cellk = ParagraphStyle("ck", parent=cell, fontName="Helvetica-Bold", textColor=c(NAVY), fontSize=8.8, leading=11)
    cellv = ParagraphStyle("cv", parent=cell, fontSize=8.8, leading=11)
    h1 = ParagraphStyle("h1", parent=base, fontName="Helvetica-Bold", fontSize=15, leading=19, textColor=c(BLUE), spaceBefore=12, spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=base, fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=c(NAVY), spaceBefore=8, spaceAfter=4)
    title = ParagraphStyle("t", parent=base, fontName="Helvetica-Bold", fontSize=26, leading=32, textColor=c(BLUE), alignment=TA_CENTER)
    sub = ParagraphStyle("s", parent=base, fontSize=15, leading=20, textColor=c(NAVY), alignment=TA_CENTER, spaceAfter=20)
    conf = ParagraphStyle("cf", parent=base, fontName="Helvetica-Bold", fontSize=8.5, leading=11, textColor=c(RED), alignment=TA_CENTER, spaceBefore=18)
    div = ParagraphStyle("d", parent=base, fontSize=9.5, textColor=c(GREY), alignment=TA_CENTER, spaceBefore=6)
    W = letter[0] - 1.5 * inch
    P = lambda text, st: Paragraph(escape(_s(text)), st)

    def table(headers, rows, kv=False):
        if not rows:
            return [P("None recorded.", base), Spacer(1, 6)]
        n = len(headers) if headers else 2
        widths = [W * 0.32, W * 0.68] if kv else None
        data = []
        if headers:
            data.append([P(h, cellh) for h in headers])
        for r in rows:
            data.append([P(v, cellk if kv and i == 0 else (cellv if kv else cell)) for i, v in enumerate(r)])
        t = Table(data, colWidths=widths or [W / n] * n, repeatRows=1 if headers else 0)
        style = [("GRID", (0, 0), (-1, -1), 0.4, c("9AA5B1")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                 ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
        if headers:
            style += [("BACKGROUND", (0, 0), (-1, 0), c(BLUE)), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, c("F5F8FB")])]
        if kv:
            style += [("BACKGROUND", (0, 0), (0, -1), c("EEF3F8"))]
        t.setStyle(TableStyle(style))
        return [t, Spacer(1, 8)]

    story = [Spacer(1, 30), Image(LOGO, width=2.6 * inch, height=2.6 * inch), Spacer(1, 10),
             P(rep["title"], title), P(rep["subtitle"], sub)]
    story += table(None, rep["meta"], kv=True)
    story += [P(CONFIDENTIAL, conf), P(DIVISION, div), PageBreak()]
    for blk in rep["blocks"]:
        t = blk["type"]
        if t == "h1":
            story.append(P(blk["text"], h1))
        elif t == "h2":
            story.append(P(blk["text"], h2))
        elif t == "p":
            story.append(P(blk["text"], base))
        elif t == "kv":
            story += table(None, blk["rows"], kv=True)
        elif t == "table":
            story += table(blk["headers"], blk["rows"])

    def decorate(canvas, doc_):
        canvas.saveState()
        if doc_.page > 1:
            canvas.setFillColor(c(BLUE))
            canvas.setFont("Helvetica-Bold", 8)
            canvas.drawString(0.75 * inch, letter[1] - 0.45 * inch, f"TEMPRUDENCE TECHNOLOGIES  |  {rep['title']}  |  {_s(rep['subtitle'])}")
            canvas.setStrokeColor(c(BLUE))
            canvas.setLineWidth(1.2)
            canvas.line(0.75 * inch, letter[1] - 0.52 * inch, letter[0] - 0.75 * inch, letter[1] - 0.52 * inch)
        canvas.setFillColor(c(GREY))
        canvas.setFont("Helvetica", 7.5)
        canvas.drawCentredString(letter[0] / 2, 0.4 * inch, FOOTER)
        canvas.drawRightString(letter[0] - 0.75 * inch, 0.4 * inch, f"Page {doc_.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=letter, leftMargin=0.75 * inch, rightMargin=0.75 * inch, topMargin=0.7 * inch,
                      bottomMargin=0.65 * inch, title=f"{rep['title']} - {_s(rep['subtitle'])}", author="Temprudence Technologies"
                      ).build(story, onFirstPage=decorate, onLaterPages=decorate)
    return buf.getvalue()


RENDERERS = {
    "docx": (render_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "odt": (render_odt, "application/vnd.oasis.opendocument.text"),
    "pdf": (render_pdf, "application/pdf"),
}
